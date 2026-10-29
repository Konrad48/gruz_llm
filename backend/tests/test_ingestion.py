from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from app.ingestion import chunk_sections, extract


def test_chunking_preserves_location_and_overlap():
    pieces = chunk_sections([("page 7", "a" * 2200)])
    assert len(pieces) == 2
    assert all(location == "page 7" for location, _ in pieces)
    assert len(pieces[0][1]) == 1800
    assert len(pieces[1][1]) == 650


def test_empty_scan_has_actionable_error():
    with pytest.raises(ValueError, match="OCR"):
        chunk_sections([("page 1", "  ")])


def test_word_tables_and_spreadsheet_locations(tmp_path: Path):
    word = Document()
    word.add_paragraph("Workshop instructions")
    word.add_table(rows=1, cols=2).cell(0, 0).text = "Valve clearance"
    doc_path = tmp_path / "manual.docx"
    word.save(doc_path)
    assert "Valve clearance" in extract(doc_path, ".docx")[0][1]
    book = Workbook()
    book.active.title = "Maintenance"
    book.active.append(["Oil", "5W-40"])
    sheet_path = tmp_path / "schedule.xlsx"
    book.save(sheet_path)
    assert extract(sheet_path, ".xlsx") == [("sheet: Maintenance", "Oil | 5W-40")]
