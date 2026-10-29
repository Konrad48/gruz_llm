from io import BytesIO

from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter
from sqlalchemy import select

from app.config import settings
from app.db import SessionLocal
from app.models import Chunk
from app.preview import document_preview


def test_word_preview_preserves_order_and_formatting(tmp_path):
    word = WordDocument()
    word.add_heading("Workshop manual", 1)
    word.add_paragraph().add_run("Check the assembly.").bold = True
    word.add_table(rows=1, cols=1).cell(0, 0).text = "Part reference"
    word.add_paragraph("Final inspection")
    path = tmp_path / "manual.docx"
    word.save(path)
    preview = document_preview(path, path.name)
    assert preview["kind"] == "word"
    blocks = preview["blocks"]
    assert blocks[0]["level"] == 1
    assert blocks[1]["runs"][0]["bold"] is True
    assert blocks[2]["rows"] == [["Part reference"]]
    assert blocks[3]["text"] == "Final inspection"


def test_spreadsheet_citation_locates_sheet_row_and_column(tmp_path):
    book = Workbook()
    book.create_sheet("Parts")
    sheet = book["Parts"]
    for row in range(1, 251):
        sheet.cell(row, 1, f"Reference {row}")
    excerpt = "Inspect the mounting plate before assembly"
    sheet.cell(151, 35, excerpt)
    path = tmp_path / "parts.xlsx"
    book.save(path)
    preview = document_preview(
        path, path.name, excerpt=excerpt, location="sheet: Parts"
    )
    assert preview["sheet"] == "Parts"
    assert preview["row_start"] == 151
    assert preview["column_start"] == 35
    assert preview["rows"][0][0] == excerpt
    assert len(preview["rows"]) == 100
    page = document_preview(path, path.name, sheet="Parts", row_start=1)
    assert page["rows"][0][0] == "Reference 1"
    assert len(page["rows"][0]) == 30


def test_csv_quoted_cells_and_pdf_target_page(tmp_path):
    csv_path = tmp_path / "parts.csv"
    csv_path.write_text('Name,Note\n"Plate, left","Inspect before assembly"\n')
    csv = document_preview(csv_path, csv_path.name)
    assert csv["rows"][1] == ["Plate, left", "Inspect before assembly"]
    pdf_path = tmp_path / "manual.pdf"
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(600, 800)
    writer.write(pdf_path)
    pdf = document_preview(pdf_path, pdf_path.name, location="page 2")
    assert pdf["page"] == 2
    assert pdf["pages"] == 3


def test_preview_source_isolation_and_original_pdf_ranges(client):
    project = client.create_project()
    docs = []
    for name in ("one.txt", "two.txt"):
        uploaded = client.post(
            f"/api/projects/{project}/documents",
            files={"file": (name, b"Check the assembly before installation.")},
        )
        docs.append(uploaded.json()["id"])
    with SessionLocal() as db:
        chunk = db.scalar(select(Chunk).where(Chunk.document_id == docs[0]))
        source = chunk.id
    preview = client.get(f"/api/documents/{docs[0]}/preview?source_id={source}")
    assert preview.status_code == 200
    assert preview.json()["excerpt"] == "Check the assembly before installation."
    assert (
        client.get(f"/api/documents/{docs[1]}/preview?source_id={source}").status_code
        == 404
    )
    assert client.get(f"/api/documents/{docs[0]}/content").status_code == 415
    writer = PdfWriter()
    writer.add_blank_page(600, 800)
    data = BytesIO()
    writer.write(data)
    uploaded = client.post(
        f"/api/projects/{project}/documents",
        files={"file": ("manual.pdf", data.getvalue(), "application/pdf")},
    )
    pdf_id = uploaded.json()["id"]
    response = client.get(f"/api/documents/{pdf_id}/content")
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("inline")
    partial = client.get(
        f"/api/documents/{pdf_id}/content", headers={"Range": "bytes=0-9"}
    )
    assert partial.status_code == 206
    assert partial.content == data.getvalue()[:10]
    from pathlib import Path

    (Path(settings.upload_dir) / pdf_id).unlink()
    assert client.get(f"/api/documents/{pdf_id}/preview").status_code == 404
