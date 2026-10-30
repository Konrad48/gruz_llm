import logging
from pathlib import Path
from zipfile import ZipFile

from docx import Document as WordDocument
from openai import AsyncOpenAI
from openpyxl import load_workbook
from pypdf import PdfReader

from .config import settings
from .db import SessionLocal
from .models import Chunk, Document
from .telemetry import get_langfuse, trace_operation

SUPPORTED = {".pdf", ".txt", ".md", ".csv", ".docx", ".xlsx"}
MAX_BYTES = 20 * 1024 * 1024
logger = logging.getLogger(__name__)


def extract(path: Path, suffix: str) -> list[tuple[str, str]]:
    if suffix in {".docx", ".xlsx"}:
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("Expanded document exceeds 80 MB.")
    if suffix == ".pdf":
        reader = PdfReader(path)
        if len(reader.pages) > 1000:
            raise ValueError("PDF exceeds 1,000 pages.")
        return [
            (f"page {i + 1}", page.extract_text() or "")
            for i, page in enumerate(reader.pages)
        ]
    if suffix == ".docx":
        doc = WordDocument(path)
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [
                " | ".join(c.text for c in row.cells)
                for table in doc.tables
                for row in table.rows
            ]
        )
        return [("document", text)]
    if suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return [
                (
                    f"sheet: {sheet.title}",
                    "\n".join(
                        " | ".join(str(v) if v is not None else "" for v in row)
                        for row in sheet.iter_rows(values_only=True)
                    ),
                )
                for sheet in book
            ]
        finally:
            book.close()
    return [("text", path.read_text(encoding="utf-8-sig"))]


def chunk_sections(
    sections: list[tuple[str, str]], size: int = 1800, overlap: int = 250
):
    chunks = []
    for location, text in sections:
        text = text.replace("\x00", "").strip()
        for start in range(0, len(text), size - overlap):
            content = text[start : start + size].strip()
            if content:
                chunks.append((location, content))
    if not chunks:
        raise ValueError("No readable text found. Scanned PDFs need OCR before upload.")
    if len(chunks) > 2000:
        raise ValueError(
            "Document exceeds 2,000 text chunks. Split it into smaller files."
        )
    return chunks


async def embed(texts: list[str]) -> list[list[float]]:
    async with AsyncOpenAI(
        api_key=settings.openai_api_key, timeout=60, max_retries=2
    ) as client:
        langfuse = get_langfuse()
        if langfuse is None:
            response = await client.embeddings.create(
                model=settings.embedding_model, input=texts, dimensions=1536
            )
        else:
            with langfuse.start_as_current_observation(
                name="openai-embeddings",
                as_type="embedding",
                model=settings.embedding_model,
                input=texts,
            ) as observation:
                response = await client.embeddings.create(
                    model=settings.embedding_model, input=texts, dimensions=1536
                )
                observation.update(
                    usage_details={"input": response.usage.prompt_tokens},
                    metadata={"dimensions": 1536, "count": len(texts)},
                )
        return [
            item.embedding
            for item in sorted(response.data, key=lambda item: item.index)
        ]


@trace_operation("index-document", capture_output=False)
async def index_document(document_id: str, path: Path, suffix: str):
    if (langfuse := get_langfuse()) is not None:
        langfuse.update_current_span(metadata={"document_id": document_id})
    try:
        from starlette.concurrency import run_in_threadpool

        sections = await run_in_threadpool(extract, path, suffix)
        pieces = chunk_sections(sections)
        vectors = []
        for offset in range(0, len(pieces), 64):
            vectors.extend(
                await embed([content for _, content in pieces[offset : offset + 64]])
            )
        with SessionLocal() as db:
            doc = db.get(Document, document_id)
            if doc is None:
                return
            for (location, content), vector in zip(pieces, vectors, strict=True):
                db.add(
                    Chunk(
                        project_id=doc.project_id,
                        document_id=doc.id,
                        location=location,
                        content=content,
                        embedding=vector,
                    )
                )
            doc.status, doc.chunk_count = "ready", len(pieces)
            db.commit()
    except Exception as exc:
        if (langfuse := get_langfuse()) is not None:
            langfuse.update_current_span(
                level="ERROR",
                status_message=str(exc),
                metadata={"document_id": document_id},
            )
        logger.exception("Document indexing failed: %s", document_id)
        with SessionLocal() as db:
            doc = db.get(Document, document_id)
            if doc:
                doc.status = "failed"
                doc.error = (
                    str(exc)[:300]
                    if isinstance(exc, ValueError)
                    else (
                        "Indexing failed. Check the API key, "
                        "provider availability, and backend logs."
                    )
                )
                db.commit()
